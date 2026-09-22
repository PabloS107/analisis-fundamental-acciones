import openpyxl

wb = openpyxl.load_workbook('data/processed/AAPL/financial_statements.xlsx', data_only=True)
print(f"Loaded workbook with {len(wb.sheetnames)} sheets.")

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    issues = []
    for r in range(1, ws.max_row + 1):
        row_label = str(ws.cell(r, 1).value or "").strip()
        for c in range(1, ws.max_column + 1):
            val = ws.cell(r, c).value
            num_fmt = ws.cell(r, c).number_format
            if val is not None:
                s_val = str(val)
                # Check for suspicious tokens
                if any(err in s_val.lower() for err in ['#ref', '#value', '#name', 'nan', 'undefined']):
                    issues.append(f"Row {r}, Col {c} ({row_label[:30]}): Suspicious value '{s_val}' (fmt: {num_fmt})")
                
                # Check for wrong multiplier format on currency or percentage metrics
                if 'x' in str(num_fmt).lower() and c > 1:
                    row_lower = row_label.lower()
                    if any(term in row_lower for term in ['$m', 'usd', 'stock', 'capital invertido', 'spread económico', 'nopat']):
                        issues.append(f"Row {r}, Col {c} ({row_label[:35]}): Has multiplier format '{num_fmt}' on currency/spread value '{s_val}'")
    
    if issues:
        print(f"\n[ISSUES IN {sheet_name}]:")
        for iss in issues[:10]:
            print(f"  - {iss}")
        if len(issues) > 10:
            print(f"  ... and {len(issues) - 10} more issues.")
    else:
        print(f"[{sheet_name}]: Clean (0 formatting anomalies detected)")
