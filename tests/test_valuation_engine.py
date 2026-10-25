"""
tests/test_valuation_engine.py - Pruebas unitarias para el motor de valoración dual:
- Reverse DCF con Expectations Investing para corporaciones (AAPL, DECK)
- Salvaguarda de FCF normalizado contra float financiero para híbridos (MELI)
- Modelo de Retorno Excedente (Excess Return Model) y P/TBV vs. ROTCE para bancos (NU)
- Generación de secciones en Word (.docx) y Markdown (.md)
"""

import sys
import json
import unittest
from pathlib import Path
import docx

scripts_dir = Path("pipeline").resolve()
sys.path.insert(0, str(scripts_dir))

import valuation_report_builder


class TestValuationEngine(unittest.TestCase):

    def setUp(self):
        self.nu_scores_path = Path("data/processed/NU/financial_scores.json")
        self.meli_scores_path = Path("data/processed/MELI/financial_scores.json")
        self.deck_scores_path = Path("data/processed/DECK/financial_scores.json")
        self.aapl_scores_path = Path("data/processed/AAPL/financial_scores.json")

    def test_01_deck_reverse_dcf(self):
        """Verifica que el Reverse DCF de DECK resuelva en el crecimiento implícito esperado."""
        if not self.deck_scores_path.exists():
            self.skipTest("Datos de DECK no encontrados.")
        with open(self.deck_scores_path, "r", encoding="utf-8") as f:
            deck_scores = json.load(f)
        vp = valuation_report_builder.get_valuation_parameters(deck_scores, "DECK")
        self.assertAlmostEqual(vp["g_implied_our_pct"], -2.0, delta=2.5)
        self.assertLess(vp["g_implied_our_pct"], 5.0)

    def test_02_aapl_reverse_dcf(self):
        """Verifica que el Reverse DCF de AAPL resuelva con precisión."""
        if not self.aapl_scores_path.exists():
            self.skipTest("Datos de AAPL no encontrados.")
        with open(self.aapl_scores_path, "r", encoding="utf-8") as f:
            aapl_scores = json.load(f)
        vp = valuation_report_builder.get_valuation_parameters(aapl_scores, "AAPL")
        self.assertFalse(vp["is_financial"])
        self.assertGreater(vp["g_implied_our_pct"], 15.0)

    def test_03_meli_valuation_parameters(self):
        """Verifica que para MELI el margen FCF y la facturación reflejen las cifras auditadas de los estados financieros."""
        if not self.meli_scores_path.exists():
            self.skipTest("Datos de MELI no encontrados.")
        with open(self.meli_scores_path, "r", encoding="utf-8") as f:
            meli_scores = json.load(f)
        vp = valuation_report_builder.get_valuation_parameters(meli_scores, "MELI")
        self.assertFalse(vp["is_financial"])
        # Margen FCF reportado de estados financieros (>= 30.0%)
        self.assertGreaterEqual(vp["fcf_margin_reported_pct"], 30.0)
        self.assertLessEqual(vp["fcf_margin_reported_pct"], 45.0)
        # Margen FCF operativo normalizado para plataformas (8.0% - 15.0%)
        self.assertGreaterEqual(vp["fcf_margin_pct"], 8.0)
        self.assertLessEqual(vp["fcf_margin_pct"], 15.0)
        # SOTP activado para PLATFORM_ECOSYSTEM
        self.assertIsNotNone(vp.get("sotp_valuation"))
        # Ingresos base derivados de canonical_summary (35.2B)
        self.assertGreaterEqual(vp["base_rev_m"], 30000.0)

    def test_04_nu_banking_excess_return_parameters(self):
        """Verifica que para NU se activen los parámetros de valoración bancaria (P/TBV y ROTCE)."""
        if not self.nu_scores_path.exists():
            self.skipTest("Datos de NU no encontrados.")
        with open(self.nu_scores_path, "r", encoding="utf-8") as f:
            nu_scores = json.load(f)
        vp = valuation_report_builder.get_valuation_parameters(nu_scores, "NU")
        self.assertTrue(vp["is_financial"])
        self.assertGreater(vp["tbv_per_share"], 1.0)
        self.assertGreaterEqual(vp["ptbv_current"], 4.5)
        self.assertGreater(vp["rotce_pct"], 20.0)
        self.assertGreater(vp["spread_economic_pct"], 10.0)
        self.assertGreater(vp["ptbv_justified"], 2.5)
        self.assertGreater(vp["fair_value_excess_return"], 3.0)

    def test_05_banking_markdown_and_docx_generation(self):
        """Verifica la generación de la sección 14 en Markdown y Word para entidades de crédito."""
        if not self.nu_scores_path.exists():
            self.skipTest("Datos de NU no encontrados.")
        with open(self.nu_scores_path, "r", encoding="utf-8") as f:
            nu_scores = json.load(f)

        md = valuation_report_builder.generate_valuation_section_md(nu_scores, "NU", "Nu Holdings Ltd.", section_num=14)
        self.assertIn("Excess Return Model", md)
        self.assertIn("P/TBV", md)
        self.assertIn("ROTCE", md)
        self.assertIn("Gordon-Shapiro bancario", md)

        doc = docx.Document()
        valuation_report_builder.add_valuation_section_docx(doc, nu_scores, "NU", "Nu Holdings Ltd.", section_num=14)
        headings = [p.text for p in doc.paragraphs if p.text]
        self.assertTrue(any("Excess Return" in h for h in headings))
        self.assertGreaterEqual(len(doc.tables), 3)

    def test_06_universal_sotp_multi_segment_meta(self):
        """Verifica que el motor SOTP Universal valore META distinguiendo Family of Apps (EBIT) y Reality Labs (I+D)."""
        meta_scores_path = Path("data/processed/META/financial_scores.json")
        if not meta_scores_path.exists():
            self.skipTest("Datos de META no encontrados.")
        with open(meta_scores_path, "r", encoding="utf-8") as f:
            meta_scores = json.load(f)
        meta_prof_path = Path("data/processed/META/company_profile.json")
        meta_prof = json.load(open(meta_prof_path, "r", encoding="utf-8")) if meta_prof_path.exists() else {}
        
        vp = valuation_report_builder.get_valuation_parameters(meta_scores, "META", profile_data=meta_prof)
        sotp = vp.get("sotp_valuation")
        self.assertIsNotNone(sotp, "META debe contar con valoración SOTP universal.")
        sb = sotp.get("segments_breakdown", [])
        self.assertGreaterEqual(len(sb), 2)
        
        foa = next((s for s in sb if "Family of Apps" in s["name"]), None)
        rl = next((s for s in sb if "Reality Labs" in s["name"]), None)
        self.assertIsNotNone(foa)
        self.assertIsNotNone(rl)
        
        # Family of Apps debe valorarse por EV/EBIT con múltiplo >= 15x
        self.assertEqual(foa["mult_type"], "EV/EBIT")
        self.assertGreaterEqual(foa["mult_val"], 15.0)
        self.assertGreater(foa["ev_segment_m"], 1_500_000.0)
        self.assertAlmostEqual(foa["op_margin_pct"], 41.0, delta=2.0)
        
        # Reality Labs debe valorarse como opción de I+D sin Enterprise Value negativo
        self.assertIn("I+D", foa["mult_type"] if "I+D" in foa["mult_type"] else rl["mult_type"])
        self.assertGreater(rl["ev_segment_m"], 0.0)
        self.assertLess(rl["op_margin_pct"], -100.0)
        
        # Conciliación de Equity y Fair Value por Acción
        self.assertGreater(sotp["fair_value_sotp"], 500.0)
        self.assertGreater(sotp["ev_total"], 2_000_000.0)

    def test_07_universal_sotp_deck_asc280(self):
        """Verifica que DECK valore HOKA y UGG por múltiplo EBIT con márgenes operativos auditados."""
        if not self.deck_scores_path.exists():
            self.skipTest("Datos de DECK no encontrados.")
        with open(self.deck_scores_path, "r", encoding="utf-8") as f:
            deck_scores = json.load(f)
        deck_prof_path = Path("data/processed/DECK/company_profile.json")
        deck_prof = json.load(open(deck_prof_path, "r", encoding="utf-8")) if deck_prof_path.exists() else {}
        
        vp = valuation_report_builder.get_valuation_parameters(deck_scores, "DECK", profile_data=deck_prof)
        sotp = vp.get("sotp_valuation")
        self.assertIsNotNone(sotp)
        sb = sotp.get("segments_breakdown", [])
        self.assertGreaterEqual(len(sb), 2)
        
        hoka = next((s for s in sb if "HOKA" in s["name"]), None)
        ugg = next((s for s in sb if "UGG" in s["name"]), None)
        self.assertIsNotNone(hoka)
        self.assertIsNotNone(ugg)
        self.assertAlmostEqual(hoka["op_margin_pct"], 34.3, delta=2.0)
        self.assertAlmostEqual(ugg["op_margin_pct"], 38.4, delta=2.0)
        self.assertEqual(hoka["mult_type"], "EV/EBIT")
        self.assertEqual(ugg["mult_type"], "EV/EBIT")

    def test_08_universal_sotp_markdown_and_docx_rendering(self):
        """Verifica que el informe Markdown y Word incluyan la tabla completa SOTP con columnas de margen y peso EV."""
        meta_scores_path = Path("data/processed/META/financial_scores.json")
        if not meta_scores_path.exists():
            self.skipTest("Datos de META no encontrados.")
        with open(meta_scores_path, "r", encoding="utf-8") as f:
            meta_scores = json.load(f)
        meta_prof_path = Path("data/processed/META/company_profile.json")
        meta_prof = json.load(open(meta_prof_path, "r", encoding="utf-8")) if meta_prof_path.exists() else {}
        
        md = valuation_report_builder.generate_valuation_section_md(meta_scores, "META", "Meta Platforms, Inc.", profile_data=meta_prof, section_num=14)
        self.assertIn("Valoración por Suma de las Partes (SOTP - Sum-of-the-Parts)", md)
        self.assertIn("Family of Apps", md)
        self.assertIn("Reality Labs", md)
        self.assertIn("Margen Op. %", md)
        self.assertIn("Enterprise Value ($ M)", md)
        
        doc = docx.Document()
        valuation_report_builder.add_valuation_section_docx(doc, meta_scores, "META", "Meta Platforms, Inc.", profile_data=meta_prof, section_num=14)
        texts = [p.text for p in doc.paragraphs if p.text]
        self.assertTrue(any("Suma de las Partes" in t for t in texts))
        self.assertGreaterEqual(len(doc.tables), 4)

    def test_09_segments_data_zero_fake_defaults(self):
        """Verifica cumplimiento de Regla 6: Cero prorratas ficticias en subsegmentos anuales de AAPL."""
        aapl_seg_path = Path("data/processed/AAPL/segments_data.json")
        if not aapl_seg_path.exists():
            self.skipTest("segments_data.json de AAPL no encontrado.")
        with open(aapl_seg_path, "r", encoding="utf-8") as f:
            data = json.load(f)
        for g in data.get("geographic_segments", []):
            if g.get("name") in ["U.S.", "China", "Other countries"]:
                self.assertTrue(g.get("is_subsegment"), f"{g['name']} debe estar clasificado como is_subsegment: True")
                self.assertNotIn("LTM", g.get("values", {}), f"{g['name']} no debe contener un LTM prorrateado artificialmente.")


if __name__ == "__main__":
    unittest.main()
