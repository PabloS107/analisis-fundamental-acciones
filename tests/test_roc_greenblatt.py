"""
tests/test_roc_greenblatt.py - Suite de pruebas unitarias y de integración
para el cálculo del Return on Capital (ROC) según la metodología Greenblatt (Magic Formula),
con Capital Tangible Operativo Promedio, salvaguardas contables institucionales y benchmark sectorial.
"""

import os
import sys
import json
import unittest
from pathlib import Path
import openpyxl

# Añadir scripts al path
scripts_dir = Path(".agents/skills/fundamental-data-extractor/scripts").resolve()
sys.path.insert(0, str(scripts_dir))

from financial_scoring_engine import (
    GreenblattROCCalculator,
    load_peer_benchmarks,
    compute_peer_benchmarks,
    calculate_fundamental_scores
)
from sec_extractor import XBRL_TAXONOMY_MAP
import excel_formatter


class TestGreenblattROC(unittest.TestCase):

    def test_01_taxonomy_map_net_ppe(self):
        """Verifica que la taxonomía XBRL incluya mapeo robusto para net_ppe en US-GAAP e IFRS."""
        self.assertIn("net_ppe", XBRL_TAXONOMY_MAP)
        concepts = XBRL_TAXONOMY_MAP["net_ppe"]
        self.assertTrue(any("PropertyPlantAndEquipmentNet" in c for c in concepts))
        self.assertTrue(any("PropertyPlantAndEquipment" in c for c in concepts))

    def test_02_calculator_positive_nwc(self):
        """Verifica el cálculo de ROC con capital de trabajo neto positivo."""
        calc = GreenblattROCCalculator()
        # Año t-1: NWC = (1000 - 200) - (500 - 100) = 800 - 400 = 400. PP&E = 600. Cap = 1000.
        # Año t: NWC = (1200 - 300) - (600 - 100) = 900 - 500 = 400. PP&E = 800. Cap = 1200.
        # Cap Promedio = (1000 + 1200) / 2 = 1100.
        # EBIT = 330 -> ROC = 330 / 1100 = 30.0%
        history = {
            "FY2023": {
                "operating_income": 250.0,
                "current_assets": 1000.0,
                "cash": 200.0,
                "current_liabilities": 500.0,
                "short_term_debt": 100.0,
                "net_ppe": 600.0
            },
            "FY2024": {
                "operating_income": 330.0,
                "current_assets": 1200.0,
                "cash": 300.0,
                "current_liabilities": 600.0,
                "short_term_debt": 100.0,
                "net_ppe": 800.0
            }
        }
        res_data = calc.calculate_series(["FY2023", "FY2024"], history, {})
        res = res_data["annual_series"]
        self.assertIn("FY2024", res)
        self.assertAlmostEqual(res["FY2024"]["tangible_capital_avg"], 1100.0, places=1)
        self.assertAlmostEqual(res["FY2024"]["roc_avg"], 0.30, places=2)

    def test_03_calculator_negative_nwc_safeguard(self):
        """Verifica la salvaguarda para modelos con capital de trabajo negativo (Asset-Light como AAPL)."""
        calc = GreenblattROCCalculator()
        # NWC = (50 - 40) - (100 - 10) = 10 - 90 = -80. Net PP&E = 50.
        # NWC + PP&E = -30. Con salvaguarda max(net_ppe, nwc + net_ppe) = max(50, -30) = 50.
        cap = calc.compute_tangible_capital(
            current_assets=50.0,
            cash=40.0,
            current_liabilities=100.0,
            short_term_debt=10.0,
            net_ppe=50.0
        )
        self.assertEqual(cap, 50.0)

    def test_04_rating_tiers(self):
        """Verifica las categorías de calificación del ROC Greenblatt."""
        calc = GreenblattROCCalculator()
        self.assertEqual(calc.rate_roc(0.55), "ELITE (> 40%)")
        self.assertEqual(calc.rate_roc(0.35), "EXCELENTE (25% - 40%)")
        self.assertEqual(calc.rate_roc(0.18), "SÓLIDO (15% - 25%)")
        self.assertEqual(calc.rate_roc(0.12), "MODERADO (10% - 15%)")
        self.assertEqual(calc.rate_roc(0.05), "DEFICIENTE (< 10%)")
        self.assertEqual(calc.rate_roc(-0.02), "DESTRUCTOR (< 0%)")

    def test_05_peer_benchmarks_contain_greenblatt_roc(self):
        """Verifica que el catálogo config/peer_benchmarks.json contenga greenblatt_roc en todos los pares."""
        raw = load_peer_benchmarks("config/peer_benchmarks.json")
        bms = raw.get("industry_benchmarks", raw)
        for sic in ["SIC_3571", "SIC_3021"]:
            self.assertIn(sic, bms)
            peers = bms[sic].get("peers", [])
            self.assertTrue(len(peers) > 0)
            for p in peers:
                self.assertIn("greenblatt_roc", p, f"Par {p.get('ticker')} en {sic} debe tener greenblatt_roc")
                self.assertIsInstance(p["greenblatt_roc"], (int, float))
            
            # Verificar compute_peer_benchmarks
            computed = compute_peer_benchmarks(bms[sic])
            self.assertIn("median_greenblatt_roc", computed)
            self.assertGreater(computed["median_greenblatt_roc"], 0.0)

    def test_06_scoring_engine_full_integration_aapl(self):
        """Verifica que calculate_fundamental_scores genere la sección greenblatt_roc_analysis completa."""
        json_path = Path("data/processed/AAPL/financial_statements.json")
        if not json_path.exists():
            self.skipTest("data/processed/AAPL/financial_statements.json no encontrado")
        
        scores = calculate_fundamental_scores("AAPL", json_path)
        self.assertIn("greenblatt_roc_analysis", scores)
        roc_data = scores["greenblatt_roc_analysis"]
        self.assertIn("annual_series", roc_data)
        self.assertIn("ltm", roc_data)
        self.assertIn("peer_comparison", roc_data)
        
        ltm_res = roc_data["ltm"]
        self.assertIsNotNone(ltm_res.get("roc_avg"))
        self.assertGreater(ltm_res["roc_avg"], 0.50, "AAPL ROC LTM debe superar 50%")
        
        peer_comp = roc_data["peer_comparison"]
        self.assertIn("peer_median_roc", peer_comp)
        self.assertIn("spread_vs_peer_median", peer_comp)
        self.assertGreater(peer_comp["spread_vs_peer_median"], 0.0)

    def test_07_excel_formatter_peer_group_sheet_col_20(self):
        """Verifica que la Hoja 8 del libro Excel incluya ROC Greenblatt en columna 20 con fórmulas dinámicas."""
        json_path = Path("data/processed/AAPL/financial_statements.json")
        if not json_path.exists():
            self.skipTest("data/processed/AAPL/financial_statements.json no encontrado")
        
        scores = calculate_fundamental_scores("AAPL", json_path)
        
        wb = openpyxl.Workbook()
        excel_formatter.create_peer_group_sheet(
            wb,
            None,
            "Apple Inc.",
            "AAPL",
            sic_code=scores.get("sic_code", "3571"),
            scoring_results=scores
        )
        ws = wb["Peer Group & Medianas"]
        
        # Header en fila 6, col 20
        self.assertEqual(ws.cell(row=6, column=20).value, "ROC Greenblatt (%)")
        
        # Encontrar fila de MEDIANA
        med_row = None
        for r in range(7, 30):
            if ws.cell(row=r, column=1).value == "MEDIANA":
                med_row = r
                break
        self.assertIsNotNone(med_row, "Fila de MEDIANA debe existir")
        val_med = ws.cell(row=med_row, column=20).value
        self.assertTrue(str(val_med).startswith("=MEDIAN(T"), f"Fórmula de mediana debe ser =MEDIAN(T...): {val_med}")

    def test_08_excel_formatter_ratios_and_scoring_sheets(self):
        """Verifica que Hoja 2 y Hoja 7 incluyan ROC Greenblatt."""
        json_path = Path("data/processed/AAPL/financial_statements.json")
        if not json_path.exists():
            self.skipTest("data/processed/AAPL/financial_statements.json no encontrado")
        
        with open(json_path, "r", encoding="utf-8") as f:
            f_json = json.load(f)
        
        fiscal_years = f_json.get("fiscal_years", ["FY2023", "FY2024", "FY2025", "LTM"])
        scores = calculate_fundamental_scores("AAPL", json_path)
        
        wb = openpyxl.Workbook()
        excel_formatter.create_ratios_sheet(wb, f_json.get("key_ratios", {}), "Apple Inc.", "AAPL", fiscal_years)
        ws_r = wb["Ratios Financieros"]
        
        # Buscar ROC en Hoja 2
        found_roc_h2 = False
        for r in range(1, 40):
            val = str(ws_r.cell(row=r, column=1).value or "")
            if "ROC Greenblatt" in val:
                found_roc_h2 = True
                break
        self.assertTrue(found_roc_h2, "ROC Greenblatt debe figurar en Hoja 2 (Ratios Financieros)")
        
        excel_formatter.create_scoring_sheet(wb, scores, "Apple Inc.", "AAPL", fiscal_years)
        ws_s = wb["Scores Fundamentales"]
        
        # Buscar ROC en Hoja 7 y verificar que LTM tenga valor numérico (no '-')
        found_roc_h7 = False
        ltm_col_idx = None
        for c in range(1, len(fiscal_years) + 5):
            if ws_s.cell(row=4, column=c).value == "LTM":
                ltm_col_idx = c
                break

        roc_h7_row = None
        for r in range(1, 120):
            val = str(ws_s.cell(row=r, column=1).value or "")
            if "ROC Greenblatt" in val:
                found_roc_h7 = True
                roc_h7_row = r
                break
        self.assertTrue(found_roc_h7, "ROC Greenblatt debe figurar en Hoja 7 (Scores Fundamentales)")
        
        if ltm_col_idx and roc_h7_row:
            ltm_roc_val = ws_s.cell(row=roc_h7_row, column=ltm_col_idx).value
            self.assertNotEqual(ltm_roc_val, "-", "ROC LTM en Hoja 7 no debe ser un guión")
            self.assertIsNotNone(ltm_roc_val, "ROC LTM en Hoja 7 no debe ser None")
            self.assertGreater(float(ltm_roc_val), 0.50, "ROC LTM en Hoja 7 debe ser mayor al 50%")

        # Verificar filas de desglose de balance para ROC
        h7_labels = [str(ws_s.cell(row=r, column=1).value or "") for r in range(1, 120)]
        self.assertTrue(any("EBIT Operativo" in l for l in h7_labels), "Debe existir fila de desglose EBIT Operativo")
        self.assertTrue(any("Capital de Trabajo Neto Operativo" in l for l in h7_labels), "Debe existir fila de desglose NWC")
        self.assertTrue(any("PP&E Neto" in l for l in h7_labels), "Debe existir fila de desglose PP&E Neto")
        self.assertTrue(any("Capital Tangible Operativo Promedio" in l for l in h7_labels), "Debe existir fila de Capital Tangible Promedio")


if __name__ == "__main__":
    unittest.main()
